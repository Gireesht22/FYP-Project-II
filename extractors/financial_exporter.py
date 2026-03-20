"""
Financial Exporter Module
Exports financial analysis results to Excel and PDF
"""

import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter


# ── COLOR PALETTE ─────────────────────────────────────────────────────────────
DARK_BG     = "0D1220"
ACCENT      = "00E5FF"
GREEN       = "43D4A0"
RED         = "FF6B6B"
ORANGE      = "FFB347"
PURPLE      = "7C5CBF"
WHITE       = "E2E8F0"
GRAY        = "445068"
LIGHT_BG    = "131929"
MID_BG      = "1A2235"

def _font(bold=False, size=11, color=WHITE, name="Arial"):
    return Font(bold=bold, size=size, color=color, name=name)

def _fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _border():
    s = Side(style="thin", color="1C2640")
    return Border(left=s, right=s, top=s, bottom=s)

def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)

def _left():
    return Alignment(horizontal="left", vertical="center", wrap_text=True)

def _money(val):
    try: return round(float(val), 2)
    except: return 0.0


# ── EXCEL EXPORT ──────────────────────────────────────────────────────────────
def export_to_excel(fin_data):
    """Generate a professional Excel report from financial analysis data"""
    wb = Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22
    ws.row_dimensions[1].height = 50

    # Title row
    ws.merge_cells("A1:D1")
    ws["A1"] = "FINANCIAL ANALYSIS REPORT"
    ws["A1"].font = _font(bold=True, size=18, color=ACCENT)
    ws["A1"].fill = _fill(DARK_BG)
    ws["A1"].alignment = _center()

    # Generated date
    ws.merge_cells("A2:D2")
    ws["A2"] = f"Generated: {datetime.now().strftime('%d %B %Y, %I:%M %p')}  |  Powered by Groq AI (Llama 3.3 70B)"
    ws["A2"].font = _font(size=9, color=GRAY)
    ws["A2"].fill = _fill(DARK_BG)
    ws["A2"].alignment = _center()
    ws.row_dimensions[2].height = 18

    # Account Info Section
    acct = fin_data.get("account_info", {})
    summary = fin_data.get("ai_summary", {})
    stats = fin_data.get("statistics", {})

    row = 4
    ws.merge_cells(f"A{row}:D{row}")
    ws[f"A{row}"] = "ACCOUNT INFORMATION"
    ws[f"A{row}"].font = _font(bold=True, size=11, color=DARK_BG)
    ws[f"A{row}"].fill = _fill(ACCENT)
    ws[f"A{row}"].alignment = _center()
    ws.row_dimensions[row].height = 22
    row += 1

    acct_fields = [
        ("Account Holder", acct.get("account_holder", "Unknown")),
        ("Bank Name", acct.get("bank_name", "Unknown")),
        ("Account Type", acct.get("account_type", "Unknown")),
        ("Account Number", acct.get("account_number", "Unknown")),
        ("Statement Period", acct.get("statement_period", "Unknown")),
        ("Currency", acct.get("currency", "INR")),
    ]
    for label, value in acct_fields:
        ws[f"A{row}"] = label
        ws[f"A{row}"].font = _font(bold=True, size=10, color=GRAY)
        ws[f"A{row}"].fill = _fill(LIGHT_BG)
        ws[f"A{row}"].border = _border()
        ws[f"A{row}"].alignment = _left()
        ws.merge_cells(f"B{row}:D{row}")
        ws[f"B{row}"] = value
        ws[f"B{row}"].font = _font(size=10, color=WHITE)
        ws[f"B{row}"].fill = _fill(MID_BG)
        ws[f"B{row}"].border = _border()
        ws[f"B{row}"].alignment = _left()
        ws.row_dimensions[row].height = 18
        row += 1

    # Financial Summary Cards
    row += 1
    ws.merge_cells(f"A{row}:D{row}")
    ws[f"A{row}"] = "FINANCIAL SUMMARY"
    ws[f"A{row}"].font = _font(bold=True, size=11, color=DARK_BG)
    ws[f"A{row}"].fill = _fill(GREEN)
    ws[f"A{row}"].alignment = _center()
    ws.row_dimensions[row].height = 22
    row += 1

    fin_fields = [
        ("Opening Balance", _money(acct.get("opening_balance", 0)), WHITE),
        ("Closing Balance", _money(acct.get("closing_balance", 0)), WHITE),
        ("Total Income (Credits)", _money(stats.get("total_credit_amount", 0)), GREEN),
        ("Total Expenses (Debits)", _money(stats.get("total_debit_amount", 0)), RED),
        ("Net Balance", _money(stats.get("net_balance", 0)), GREEN if _money(stats.get("net_balance",0)) >= 0 else RED),
        ("Savings Rate", f"{stats.get('savings_rate', 0)}%", GREEN),
        ("Average Credit", _money(stats.get("avg_credit", 0)), WHITE),
        ("Average Debit", _money(stats.get("avg_debit", 0)), WHITE),
        ("Largest Credit", _money(stats.get("largest_credit", 0)), GREEN),
        ("Largest Debit", _money(stats.get("largest_debit", 0)), RED),
        ("Total Transactions", stats.get("total_transactions", 0), WHITE),
    ]
    for label, value, color in fin_fields:
        ws[f"A{row}"] = label
        ws[f"A{row}"].font = _font(bold=True, size=10, color=GRAY)
        ws[f"A{row}"].fill = _fill(LIGHT_BG)
        ws[f"A{row}"].border = _border()
        ws[f"A{row}"].alignment = _left()
        ws.merge_cells(f"B{row}:D{row}")
        ws[f"B{row}"] = value
        ws[f"B{row}"].font = _font(bold=True, size=11, color=color)
        ws[f"B{row}"].fill = _fill(MID_BG)
        ws[f"B{row}"].border = _border()
        ws[f"B{row}"].alignment = _left()
        if isinstance(value, float):
            ws[f"B{row}"].number_format = '#,##0.00'
        ws.row_dimensions[row].height = 18
        row += 1

    # AI Health Score
    row += 1
    health = summary.get("financial_health", "Unknown")
    score = summary.get("health_score", 0)
    hcolor = GREEN if health == "Good" else ORANGE if health == "Fair" else RED
    ws.merge_cells(f"A{row}:D{row}")
    ws[f"A{row}"] = "AI FINANCIAL HEALTH"
    ws[f"A{row}"].font = _font(bold=True, size=11, color=DARK_BG)
    ws[f"A{row}"].fill = _fill(PURPLE)
    ws[f"A{row}"].alignment = _center()
    ws.row_dimensions[row].height = 22
    row += 1

    ws[f"A{row}"] = "Health Rating"
    ws[f"A{row}"].font = _font(bold=True, size=10, color=GRAY)
    ws[f"A{row}"].fill = _fill(LIGHT_BG)
    ws[f"A{row}"].border = _border()
    ws.merge_cells(f"B{row}:D{row}")
    ws[f"B{row}"] = f"{health} ({score}/100)"
    ws[f"B{row}"].font = _font(bold=True, size=12, color=hcolor)
    ws[f"B{row}"].fill = _fill(MID_BG)
    ws[f"B{row}"].border = _border()
    row += 1

    ws[f"A{row}"] = "Executive Summary"
    ws[f"A{row}"].font = _font(bold=True, size=10, color=GRAY)
    ws[f"A{row}"].fill = _fill(LIGHT_BG)
    ws[f"A{row}"].border = _border()
    ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="top")
    ws.merge_cells(f"B{row}:D{row}")
    ws[f"B{row}"] = summary.get("executive_summary", "")
    ws[f"B{row}"].font = _font(size=10, color=WHITE)
    ws[f"B{row}"].fill = _fill(MID_BG)
    ws[f"B{row}"].border = _border()
    ws[f"B{row}"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[row].height = 50
    row += 1

    # Key Insights
    insights = summary.get("key_insights", [])
    if insights:
        ws[f"A{row}"] = "Key Insights"
        ws[f"A{row}"].font = _font(bold=True, size=10, color=GRAY)
        ws[f"A{row}"].fill = _fill(LIGHT_BG)
        ws[f"A{row}"].border = _border()
        ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="top")
        ws.merge_cells(f"B{row}:D{row}")
        ws[f"B{row}"] = "\n".join([f"• {i}" for i in insights])
        ws[f"B{row}"].font = _font(size=10, color=WHITE)
        ws[f"B{row}"].fill = _fill(MID_BG)
        ws[f"B{row}"].border = _border()
        ws[f"B{row}"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = max(60, len(insights) * 18)
        row += 1

    # Recommendations
    recs = summary.get("recommendations", [])
    if recs:
        ws[f"A{row}"] = "Recommendations"
        ws[f"A{row}"].font = _font(bold=True, size=10, color=GRAY)
        ws[f"A{row}"].fill = _fill(LIGHT_BG)
        ws[f"A{row}"].border = _border()
        ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="top")
        ws.merge_cells(f"B{row}:D{row}")
        ws[f"B{row}"] = "\n".join([f"→ {r}" for r in recs])
        ws[f"B{row}"].font = _font(size=10, color=ORANGE)
        ws[f"B{row}"].fill = _fill(MID_BG)
        ws[f"B{row}"].border = _border()
        ws[f"B{row}"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        ws.row_dimensions[row].height = max(60, len(recs) * 18)
        row += 1

    # ── Sheet 2: Transactions ─────────────────────────────────────────────────
    ws2 = wb.create_sheet("Transactions")
    ws2.sheet_view.showGridLines = False
    transactions = fin_data.get("transactions", [])

    col_widths = [15, 40, 12, 12, 18, 18]
    col_letters = ["A","B","C","D","E","F"]
    for i, w in enumerate(col_widths):
        ws2.column_dimensions[col_letters[i]].width = w

    # Header
    ws2.merge_cells("A1:F1")
    ws2["A1"] = "TRANSACTION DETAILS"
    ws2["A1"].font = _font(bold=True, size=14, color=ACCENT)
    ws2["A1"].fill = _fill(DARK_BG)
    ws2["A1"].alignment = _center()
    ws2.row_dimensions[1].height = 35

    headers = ["Date", "Description", "Type", "Category", "Amount", "Balance"]
    for i, h in enumerate(headers):
        cell = ws2.cell(row=2, column=i+1, value=h)
        cell.font = _font(bold=True, size=10, color=DARK_BG)
        cell.fill = _fill(ACCENT)
        cell.alignment = _center()
        cell.border = _border()
    ws2.row_dimensions[2].height = 20

    for r, t in enumerate(transactions, start=3):
        row_data = [
            t.get("date", ""),
            t.get("description", ""),
            t.get("type", ""),
            t.get("category", ""),
            _money(t.get("amount", 0)),
            _money(t.get("balance", 0)),
        ]
        is_credit = t.get("type") == "Credit"
        row_fill = LIGHT_BG if r % 2 == 0 else MID_BG
        for c, val in enumerate(row_data, start=1):
            cell = ws2.cell(row=r, column=c, value=val)
            cell.fill = _fill(row_fill)
            cell.border = _border()
            cell.alignment = _left()
            if c == 5:  # Amount column
                cell.font = _font(bold=True, size=10, color=GREEN if is_credit else RED)
                cell.number_format = '#,##0.00'
            elif c == 6:  # Balance
                cell.font = _font(size=10, color=WHITE)
                cell.number_format = '#,##0.00'
            elif c == 3:  # Type
                cell.font = _font(bold=True, size=10, color=GREEN if is_credit else RED)
            else:
                cell.font = _font(size=10, color=WHITE)
        ws2.row_dimensions[r].height = 16

    # Total row
    if transactions:
        total_row = len(transactions) + 3
        ws2.merge_cells(f"A{total_row}:D{total_row}")
        ws2[f"A{total_row}"] = "TOTALS"
        ws2[f"A{total_row}"].font = _font(bold=True, size=11, color=DARK_BG)
        ws2[f"A{total_row}"].fill = _fill(ACCENT)
        ws2[f"A{total_row}"].alignment = _center()
        ws2[f"E{total_row}"] = f"=SUMIF(C3:C{total_row-1},\"Credit\",E3:E{total_row-1})-SUMIF(C3:C{total_row-1},\"Debit\",E3:E{total_row-1})"
        ws2[f"E{total_row}"].font = _font(bold=True, size=11, color=GREEN)
        ws2[f"E{total_row}"].fill = _fill(DARK_BG)
        ws2[f"E{total_row}"].number_format = '#,##0.00'
        ws2[f"E{total_row}"].border = _border()

    # ── Sheet 3: Category Analysis ────────────────────────────────────────────
    ws3 = wb.create_sheet("Category Analysis")
    ws3.sheet_view.showGridLines = False
    ws3.column_dimensions["A"].width = 25
    ws3.column_dimensions["B"].width = 20
    ws3.column_dimensions["C"].width = 15

    ws3.merge_cells("A1:C1")
    ws3["A1"] = "SPENDING BY CATEGORY"
    ws3["A1"].font = _font(bold=True, size=14, color=ACCENT)
    ws3["A1"].fill = _fill(DARK_BG)
    ws3["A1"].alignment = _center()
    ws3.row_dimensions[1].height = 35

    cat_headers = ["Category", "Amount", "% of Total"]
    for i, h in enumerate(cat_headers):
        cell = ws3.cell(row=2, column=i+1, value=h)
        cell.font = _font(bold=True, size=10, color=DARK_BG)
        cell.fill = _fill(PURPLE)
        cell.alignment = _center()
        cell.border = _border()
    ws3.row_dimensions[2].height = 20

    cats = fin_data.get("statistics", {}).get("category_breakdown", {})
    total_spend = sum(cats.values()) or 1
    cat_colors = [ACCENT, PURPLE, GREEN, ORANGE, RED, "A78BFA", "34D399", "F472B6"]
    for r, (cat, amt) in enumerate(sorted(cats.items(), key=lambda x: -x[1]), start=3):
        pct = round(amt / total_spend * 100, 1)
        color = cat_colors[(r-3) % len(cat_colors)]
        for c, val in enumerate([cat, _money(amt), f"{pct}%"], start=1):
            cell = ws3.cell(row=r, column=c, value=val)
            cell.font = _font(size=10, color=color if c == 1 else WHITE)
            cell.fill = _fill(LIGHT_BG if r % 2 == 0 else MID_BG)
            cell.border = _border()
            cell.alignment = _left()
            if c == 2:
                cell.number_format = '#,##0.00'
        ws3.row_dimensions[r].height = 18

    # Total row
    last_cat_row = len(cats) + 3
    ws3[f"A{last_cat_row}"] = "TOTAL"
    ws3[f"A{last_cat_row}"].font = _font(bold=True, size=11, color=DARK_BG)
    ws3[f"A{last_cat_row}"].fill = _fill(GREEN)
    ws3[f"A{last_cat_row}"].border = _border()
    ws3[f"B{last_cat_row}"] = f"=SUM(B3:B{last_cat_row-1})"
    ws3[f"B{last_cat_row}"].font = _font(bold=True, size=11, color=DARK_BG)
    ws3[f"B{last_cat_row}"].fill = _fill(GREEN)
    ws3[f"B{last_cat_row}"].border = _border()
    ws3[f"B{last_cat_row}"].number_format = '#,##0.00'
    ws3[f"C{last_cat_row}"] = "100%"
    ws3[f"C{last_cat_row}"].font = _font(bold=True, size=11, color=DARK_BG)
    ws3[f"C{last_cat_row}"].fill = _fill(GREEN)
    ws3[f"C{last_cat_row}"].border = _border()

    # ── Sheet 4: Risk Flags ───────────────────────────────────────────────────
    ws4 = wb.create_sheet("Risk Flags")
    ws4.sheet_view.showGridLines = False
    ws4.column_dimensions["A"].width = 40
    ws4.column_dimensions["B"].width = 18
    ws4.column_dimensions["C"].width = 45

    ws4.merge_cells("A1:C1")
    ws4["A1"] = "RISK FLAGS & SUSPICIOUS TRANSACTIONS"
    ws4["A1"].font = _font(bold=True, size=14, color=RED)
    ws4["A1"].fill = _fill(DARK_BG)
    ws4["A1"].alignment = _center()
    ws4.row_dimensions[1].height = 35

    risk_headers = ["Transaction", "Amount", "Reason"]
    for i, h in enumerate(risk_headers):
        cell = ws4.cell(row=2, column=i+1, value=h)
        cell.font = _font(bold=True, size=10, color=DARK_BG)
        cell.fill = _fill(RED)
        cell.alignment = _center()
        cell.border = _border()
    ws4.row_dimensions[2].height = 20

    flags = fin_data.get("statistics", {}).get("risk_flags", [])
    if flags:
        for r, flag in enumerate(flags, start=3):
            for c, val in enumerate([flag.get("transaction",""), _money(flag.get("amount",0)), flag.get("reason","")], start=1):
                cell = ws4.cell(row=r, column=c, value=val)
                cell.font = _font(size=10, color=RED if c == 2 else WHITE)
                cell.fill = _fill(LIGHT_BG)
                cell.border = _border()
                cell.alignment = _left()
                if c == 2: cell.number_format = '#,##0.00'
            ws4.row_dimensions[r].height = 18
    else:
        ws4.merge_cells("A3:C3")
        ws4["A3"] = "✅ No suspicious transactions detected"
        ws4["A3"].font = _font(size=12, color=GREEN)
        ws4["A3"].fill = _fill(LIGHT_BG)
        ws4["A3"].alignment = _center()

    output = io.BytesIO()
    wb.save(output)
    return output


# ── PDF EXPORT ────────────────────────────────────────────────────────────────
def export_to_pdf(fin_data):
    """Generate a professional PDF report from financial analysis data"""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, HRFlowable)
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4,
                            leftMargin=15*mm, rightMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm)

    # Colors
    C_DARK   = colors.HexColor("#0D1220")
    C_ACCENT = colors.HexColor("#00E5FF")
    C_GREEN  = colors.HexColor("#43D4A0")
    C_RED    = colors.HexColor("#FF6B6B")
    C_ORANGE = colors.HexColor("#FFB347")
    C_PURPLE = colors.HexColor("#7C5CBF")
    C_WHITE  = colors.HexColor("#E2E8F0")
    C_GRAY   = colors.HexColor("#8B9AB5")
    C_BG2    = colors.HexColor("#131929")
    C_BG3    = colors.HexColor("#1A2235")

    styles = getSampleStyleSheet()
    def sty(name, **kwargs):
        return ParagraphStyle(name, **kwargs)

    title_sty  = sty("Title2", fontSize=22, textColor=C_ACCENT, alignment=TA_CENTER, fontName="Helvetica-Bold", spaceAfter=4)
    sub_sty    = sty("Sub",    fontSize=9,  textColor=C_GRAY,   alignment=TA_CENTER, fontName="Helvetica", spaceAfter=12)
    sec_sty    = sty("Sec",    fontSize=12, textColor=C_DARK,   alignment=TA_CENTER, fontName="Helvetica-Bold", spaceAfter=6, spaceBefore=12, backColor=C_ACCENT, borderPadding=6)
    body_sty   = sty("Body2",  fontSize=10, textColor=C_WHITE,  fontName="Helvetica", spaceAfter=4, leading=15)
    label_sty  = sty("Label",  fontSize=9,  textColor=C_GRAY,   fontName="Helvetica-Bold", spaceAfter=2)
    green_sty  = sty("Green",  fontSize=11, textColor=C_GREEN,  fontName="Helvetica-Bold")
    red_sty    = sty("Red",    fontSize=11, textColor=C_RED,    fontName="Helvetica-Bold")
    orange_sty = sty("Orange", fontSize=10, textColor=C_ORANGE, fontName="Helvetica")

    story = []
    acct  = fin_data.get("account_info", {})
    stats = fin_data.get("statistics", {})
    summary = fin_data.get("ai_summary", {})
    transactions = fin_data.get("transactions", [])

    # ── Title ─────────────────────────────────────────────────────────────────
    story.append(Paragraph("FINANCIAL ANALYSIS REPORT", title_sty))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%d %B %Y, %I:%M %p')}  |  Powered by Groq AI", sub_sty))
    story.append(HRFlowable(width="100%", thickness=1, color=C_ACCENT))
    story.append(Spacer(1, 8))

    # ── Account Info ──────────────────────────────────────────────────────────
    story.append(Paragraph("ACCOUNT INFORMATION", sec_sty))
    acct_data = [
        ["Account Holder", acct.get("account_holder","Unknown"), "Bank Name", acct.get("bank_name","Unknown")],
        ["Account Type",   acct.get("account_type","Unknown"),   "Currency",  acct.get("currency","INR")],
        ["Account Number", acct.get("account_number","Unknown"), "Period",    acct.get("statement_period","Unknown")],
    ]
    acct_table = Table(acct_data, colWidths=[40*mm, 55*mm, 35*mm, 55*mm])
    acct_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), C_BG3),
        ("TEXTCOLOR", (0,0), (0,-1), C_GRAY),
        ("TEXTCOLOR", (2,0), (2,-1), C_GRAY),
        ("TEXTCOLOR", (1,0), (1,-1), C_WHITE),
        ("TEXTCOLOR", (3,0), (3,-1), C_WHITE),
        ("FONTNAME", (0,0), (0,-1), "Helvetica-Bold"),
        ("FONTNAME", (2,0), (2,-1), "Helvetica-Bold"),
        ("FONTNAME", (1,0), (1,-1), "Helvetica"),
        ("FONTNAME", (3,0), (3,-1), "Helvetica"),
        ("FONTSIZE", (0,0), (-1,-1), 9),
        ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor("#1C2640")),
        ("ROWBACKGROUNDS", (0,0), (-1,-1), [C_BG2, C_BG3]),
        ("PADDING", (0,0), (-1,-1), 7),
    ]))
    story.append(acct_table)
    story.append(Spacer(1, 8))

    # ── Financial Summary Cards ────────────────────────────────────────────────
    story.append(Paragraph("FINANCIAL SUMMARY", sec_sty))
    net = stats.get("net_balance", 0)
    summary_data = [
        ["Total Income", f"{stats.get('total_credit_amount',0):,.2f}", "Total Expenses", f"{stats.get('total_debit_amount',0):,.2f}"],
        ["Net Balance",  f"{net:,.2f}",                                "Savings Rate",   f"{stats.get('savings_rate',0)}%"],
        ["Avg Credit",   f"{stats.get('avg_credit',0):,.2f}",          "Avg Debit",      f"{stats.get('avg_debit',0):,.2f}"],
        ["Largest Credit",f"{stats.get('largest_credit',0):,.2f}",     "Largest Debit",  f"{stats.get('largest_debit',0):,.2f}"],
        ["Total Txns",   str(stats.get("total_transactions",0)),        "Risk Flags",     str(len(stats.get("risk_flags",[])))],
    ]
    sum_table = Table(summary_data, colWidths=[40*mm, 55*mm, 40*mm, 50*mm])
    sum_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), C_BG3),
        ("TEXTCOLOR", (0,0), (0,-1), C_GRAY),
        ("TEXTCOLOR", (2,0), (2,-1), C_GRAY),
        ("TEXTCOLOR", (1,0), (1,-1), C_GREEN),
        ("TEXTCOLOR", (3,0), (3,-1), C_RED),
        ("TEXTCOLOR", (1,1), (1,1), C_GREEN if net >= 0 else C_RED),
        ("FONTNAME", (0,0), (-1,-1), "Helvetica"),
        ("FONTNAME", (0,0), (0,-1), "Helvetica-Bold"),
        ("FONTNAME", (2,0), (2,-1), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 10),
        ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor("#1C2640")),
        ("ROWBACKGROUNDS", (0,0), (-1,-1), [C_BG2, C_BG3]),
        ("PADDING", (0,0), (-1,-1), 8),
        ("ALIGN", (1,0), (1,-1), "RIGHT"),
        ("ALIGN", (3,0), (3,-1), "RIGHT"),
    ]))
    story.append(sum_table)
    story.append(Spacer(1, 8))

    # ── AI Health Score ────────────────────────────────────────────────────────
    health = summary.get("financial_health","Unknown")
    score  = summary.get("health_score", 0)
    hcolor = C_GREEN if health=="Good" else C_ORANGE if health=="Fair" else C_RED
    story.append(Paragraph("AI FINANCIAL HEALTH ASSESSMENT", sec_sty))
    health_data = [
        ["Health Rating", f"{health}  ({score}/100)", "Risk Level", summary.get("risk_assessment","Unknown")],
    ]
    health_table = Table(health_data, colWidths=[40*mm, 65*mm, 35*mm, 45*mm])
    health_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), C_BG2),
        ("TEXTCOLOR", (0,0), (0,0), C_GRAY),
        ("TEXTCOLOR", (1,0), (1,0), hcolor),
        ("TEXTCOLOR", (2,0), (2,0), C_GRAY),
        ("TEXTCOLOR", (3,0), (3,0), C_RED),
        ("FONTNAME", (0,0), (-1,-1), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 11),
        ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor("#1C2640")),
        ("PADDING", (0,0), (-1,-1), 10),
    ]))
    story.append(health_table)
    story.append(Spacer(1, 6))

    exec_sum = summary.get("executive_summary","")
    if exec_sum:
        story.append(Paragraph("Executive Summary:", label_sty))
        story.append(Paragraph(exec_sum, body_sty))
        story.append(Spacer(1, 4))

    insights = summary.get("key_insights", [])
    if insights:
        story.append(Paragraph("Key Insights:", label_sty))
        for ins in insights:
            story.append(Paragraph(f"• {ins}", body_sty))
        story.append(Spacer(1, 4))

    recs = summary.get("recommendations", [])
    if recs:
        story.append(Paragraph("Recommendations:", label_sty))
        for rec in recs:
            story.append(Paragraph(f"→ {rec}", orange_sty))
        story.append(Spacer(1, 8))

    # ── Category Breakdown ────────────────────────────────────────────────────
    cats = stats.get("category_breakdown", {})
    if cats:
        story.append(Paragraph("SPENDING BY CATEGORY", sec_sty))
        total_spend = sum(cats.values()) or 1
        cat_data = [["Category", "Amount", "% of Total"]]
        for cat, amt in sorted(cats.items(), key=lambda x: -x[1]):
            cat_data.append([cat, f"{amt:,.2f}", f"{amt/total_spend*100:.1f}%"])
        cat_data.append(["TOTAL", f"{total_spend:,.2f}", "100%"])
        cat_table = Table(cat_data, colWidths=[80*mm, 55*mm, 50*mm])
        cat_table.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), C_PURPLE),
            ("BACKGROUND", (0,-1), (-1,-1), C_GREEN),
            ("TEXTCOLOR", (0,0), (-1,0), C_DARK),
            ("TEXTCOLOR", (0,-1), (-1,-1), C_DARK),
            ("TEXTCOLOR", (0,1), (-1,-2), C_WHITE),
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTNAME", (0,-1), (-1,-1), "Helvetica-Bold"),
            ("FONTNAME", (0,1), (-1,-2), "Helvetica"),
            ("FONTSIZE", (0,0), (-1,-1), 9),
            ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor("#1C2640")),
            ("ROWBACKGROUNDS", (0,1), (-1,-2), [C_BG2, C_BG3]),
            ("ALIGN", (1,0), (-1,-1), "RIGHT"),
            ("PADDING", (0,0), (-1,-1), 7),
        ]))
        story.append(cat_table)
        story.append(Spacer(1, 8))

    # ── Risk Flags ────────────────────────────────────────────────────────────
    flags = stats.get("risk_flags", [])
    story.append(Paragraph("RISK FLAGS", sec_sty))
    if flags:
        risk_data = [["Transaction", "Amount", "Reason"]]
        for f in flags:
            risk_data.append([f.get("transaction","")[:40], f"{_money(f.get('amount',0)):,.2f}", f.get("reason","")[:50]])
        risk_table = Table(risk_data, colWidths=[65*mm, 30*mm, 90*mm])
        risk_table.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), C_RED),
            ("TEXTCOLOR", (0,0), (-1,0), C_DARK),
            ("TEXTCOLOR", (0,1), (-1,-1), C_WHITE),
            ("TEXTCOLOR", (1,1), (1,-1), C_RED),
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTNAME", (0,1), (-1,-1), "Helvetica"),
            ("FONTSIZE", (0,0), (-1,-1), 9),
            ("GRID", (0,0), (-1,-1), 0.5, colors.HexColor("#1C2640")),
            ("ROWBACKGROUNDS", (0,1), (-1,-1), [C_BG2, C_BG3]),
            ("PADDING", (0,0), (-1,-1), 7),
        ]))
        story.append(risk_table)
    else:
        story.append(Paragraph("✅ No suspicious transactions detected", green_sty))
    story.append(Spacer(1, 8))

    # ── Transactions ──────────────────────────────────────────────────────────
    if transactions:
        story.append(Paragraph("TRANSACTION DETAILS", sec_sty))
        txn_data = [["Date", "Description", "Type", "Category", "Amount"]]
        for t in transactions[:50]:  # limit to 50 for PDF
            txn_data.append([
                t.get("date",""),
                t.get("description","")[:35],
                t.get("type",""),
                t.get("category",""),
                f"{_money(t.get('amount',0)):,.2f}",
            ])
        txn_table = Table(txn_data, colWidths=[22*mm, 65*mm, 20*mm, 25*mm, 28*mm])
        txn_style = [
            ("BACKGROUND", (0,0), (-1,0), C_ACCENT),
            ("TEXTCOLOR", (0,0), (-1,0), C_DARK),
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTNAME", (0,1), (-1,-1), "Helvetica"),
            ("FONTSIZE", (0,0), (-1,-1), 8),
            ("GRID", (0,0), (-1,-1), 0.3, colors.HexColor("#1C2640")),
            ("ROWBACKGROUNDS", (0,1), (-1,-1), [C_BG2, C_BG3]),
            ("ALIGN", (4,0), (4,-1), "RIGHT"),
            ("PADDING", (0,0), (-1,-1), 5),
        ]
        for i, t in enumerate(transactions[:50], start=1):
            color = C_GREEN if t.get("type") == "Credit" else C_RED
            txn_style.append(("TEXTCOLOR", (2,i), (2,i), color))
            txn_style.append(("TEXTCOLOR", (4,i), (4,i), color))
            txn_style.append(("FONTNAME", (2,i), (2,i), "Helvetica-Bold"))
        txn_table.setStyle(TableStyle(txn_style))
        story.append(txn_table)
        if len(transactions) > 50:
            story.append(Spacer(1,4))
            story.append(Paragraph(f"* Showing 50 of {len(transactions)} transactions. Full list available in Excel export.", label_sty))

    # ── Footer ────────────────────────────────────────────────────────────────
    story.append(Spacer(1, 12))
    story.append(HRFlowable(width="100%", thickness=0.5, color=C_GRAY))
    story.append(Spacer(1, 4))
    story.append(Paragraph("UniExtract AI · FYP Semester 8 · Powered by Groq AI (Llama 3.3 70B) · Privacy-First Local Processing", sub_sty))

    doc.build(story)
    return output
